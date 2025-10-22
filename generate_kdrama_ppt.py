import os
import requests
import tempfile
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt

OUTPUT = Path(r"c:\work\k-drama.pptx")
TMDB_API_KEY = os.environ.get("TMDB_API_KEY")  # 필수: TMDb API 키를 환경변수로 설정

TMDB_DISCOVER_URL = "https://api.themoviedb.org/3/discover/tv"
TMDB_IMAGE_BASE = "https://image.tmdb.org/t/p/w500"

def _tmdb_get(url, params, timeout=10):
    params = params.copy()
    params["api_key"] = TMDB_API_KEY
    r = requests.get(url, params=params, timeout=timeout)
    r.raise_for_status()
    return r.json()

def fetch_popular_kdramas(limit=6, region="KR", language="ko-KR"):
    if not TMDB_API_KEY:
        raise RuntimeError("TMDB_API_KEY 환경변수가 설정되어 있지 않습니다.")
    params = {
        "sort_by": "popularity.desc",
        "with_original_language": "ko",
        "page": 1,
        "language": language,
    }
    results = []
    while len(results) < limit:
        data = _tmdb_get(TMDB_DISCOVER_URL, params)
        for item in data.get("results", []):
            results.append({
                "title": item.get("name") or item.get("original_name"),
                "overview": item.get("overview") or "",
                "poster_path": (TMDB_IMAGE_BASE + item["poster_path"]) if item.get("poster_path") else None,
            })
            if len(results) >= limit:
                break
        params["page"] = params.get("page", 1) + 1
        if params["page"] > data.get("total_pages", 1):
            break
    return results[:limit]

def download_image(url):
    if not url:
        return None
    r = requests.get(url, stream=True, timeout=10)
    r.raise_for_status()
    suffix = os.path.splitext(url)[1].split("?")[0] or ".jpg"
    fd, path = tempfile.mkstemp(suffix=suffix)
    os.close(fd)
    with open(path, "wb") as f:
        for chunk in r.iter_content(1024):
            if not chunk:
                break
            f.write(chunk)
    return path

def create_ppt(items, output_path=OUTPUT):
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]  # blank
    for item in items:
        slide = prs.slides.add_slide(blank_layout)

        # 제목
        left = Inches(0.5)
        top = Inches(0.2)
        width = Inches(9)
        height = Inches(0.9)
        title_box = slide.shapes.add_textbox(left, top, width, height)
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = item.get("title", "제목 없음")
        p.font.size = Pt(28)
        p.font.bold = True

        # 요약
        left = Inches(0.5)
        top = Inches(1.3)
        width = Inches(6.2)
        height = Inches(4.2)
        body = slide.shapes.add_textbox(left, top, width, height)
        tf = body.text_frame
        tf.word_wrap = True
        text = item.get("overview") or "요약 없음"
        if len(text) > 900:
            text = text[:900].rsplit(" ", 1)[0] + "…"
        tf.text = text
        for paragraph in tf.paragraphs:
            paragraph.font.size = Pt(12)

        # 이미지(오른쪽)
        img_path = item.get("image_path")
        if img_path:
            try:
                img_left = Inches(6.7)
                img_top = Inches(1.3)
                img_width = Inches(3)
                slide.shapes.add_picture(img_path, img_left, img_top, width=img_width)
            except Exception:
                pass

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    print(f"저장됨: {output_path}")

def build_kdrama_ppt(limit=6):
    titles = fetch_popular_kdramas(limit=limit)
    items = []
    temp_files = []
    for t in titles:
        img_path = None
        if t.get("poster_path"):
            try:
                img_path = download_image(t["poster_path"])
                temp_files.append(img_path)
            except Exception:
                img_path = None
        items.append({
            "title": t.get("title"),
            "overview": t.get("overview"),
            "image_path": img_path,
        })
    try:
        create_ppt(items)
    finally:
        for f in temp_files:
            try:
                os.remove(f)
            except Exception:
                pass

if __name__ == "__main__":
    # 기본 실행: 인기 한국 드라마 6개로 PPT 생성
    build_kdrama_ppt(limit=6)