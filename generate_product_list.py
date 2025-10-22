import random
from pathlib import Path
from datetime import datetime

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
except Exception as e:
    raise ImportError("openpyxl이 필요합니다. 'pip install -r requirements.txt' 또는 'pip install openpyxl'로 설치하세요.") from e

OUTPUT_PATH = Path(r"c:\work\ProductList.xlsx")

BRANDS = ["Samson", "Pixelon", "Elexa", "Novatech", "Orion", "Lumix", "ZenCore", "Astra", "Vortex", "Nexon"]
TYPES = ["스마트폰", "노트북", "태블릿", "무선이어폰", "스마트워치", "TV", "블루투스스피커", "게임기", "카메라", "모니터"]

def generate_products(count=100):
    products = []
    for i in range(1, count + 1):
        product_id = f"PROD{i:04d}"
        name = f"{random.choice(BRANDS)} {random.choice(TYPES)} {random.randint(1,99)}"
        price = random.randint(30_000, 2_000_000)  # 가격 (원 단위)
        quantity = random.randint(0, 200)
        products.append((product_id, name, price, quantity))
    return products

def save_to_excel(products, path=OUTPUT_PATH):
    wb = Workbook()
    ws = wb.active
    ws.title = "ProductList"

    headers = ["제품ID", "제품명", "가격", "수량"]
    ws.append(headers)
    for row in products:
        ws.append(row)

    # 스타일: 헤더 굵게, 열 정렬 및 너비 조정
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    column_widths = [12, 40, 12, 8]
    for i, width in enumerate(column_widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width

    # 파일 저장 (덮어쓰기)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)

if __name__ == "__main__":
    products = generate_products(100)
    save_to_excel(products)
    print(f"생성 완료: {OUTPUT_PATH} ({len(products)}개, 생성시각: {datetime.now().isoformat()})")